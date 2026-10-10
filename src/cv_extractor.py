import os
import re

from pypdf import PdfReader
from docx import Document


class CVExtractionError(Exception):
    pass


def clean_text(text):
    text = text.replace("\x00", "")                
    text = text.replace("\u00a0", " ")                
    text = re.sub(r"[ \t]+", " ", text)          
    text = re.sub(r" *\n *", "\n", text)            
    text = re.sub(r"\n{3,}", "\n\n", text)          
    return text.strip()


def extract_text_from_pdf(path):
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise CVExtractionError("PDF is password protected")
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]

    
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_cv_text(path):
    
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            raw = extract_text_from_pdf(path)
        elif ext == ".docx":
            raw = extract_text_from_docx(path)
        else:
            raise CVExtractionError(f"Unsupported file type: {ext}")
    except CVExtractionError:
        raise
    except Exception as e:
        raise CVExtractionError(f"Could not read file: {e}")

    text = clean_text(raw)
    if len(text) < 50:
        raise CVExtractionError(
            "No readable text found (the file may be a scanned image)"
        )
    return text


def build_candidate(email, path):
    
    return {
        "email_id": email["id"],
        "sender": email["sender"],
        "subject": email["subject"],
        "date": email["date"],
        "cv_file": os.path.basename(path),
        "cv_text": extract_cv_text(path),
    }